"""Workbook semantic review tests with source-free synthetic XLSX fixtures."""
from pathlib import Path

from openpyxl import Workbook
import pytest

from residual_pathways.metadata_sheet_review import review, SheetReviewError


def make(path, identity_rows, *, readme_note="original", illustrative=("alpha","beta")):
    wb=Workbook()
    meta=wb.active
    meta.title="meta"
    for row in identity_rows:
        meta.append(row)
    vet=wb.create_sheet("meta_Ch3vetted_withclimate")
    for row in identity_rows:
        vet.append(row)
    readme=wb.create_sheet("README")
    readme["A1"]="Text"
    readme["A2"]=readme_note
    illustration=wb.create_sheet("Illustrative_Pathways")
    illustration.append(["Example"])
    for word in illustrative:
        illustration.append([word])
    wb.save(path)


def test_joined_records_equal_despite_row_reordering(tmp_path):
    first=tmp_path/"a.xlsx"
    second=tmp_path/"b.xlsx"
    header=["Model","Scenario","Category"]
    a=[header,["model-A","path-1","C1"],["model-B","path-2","C2"]]
    b=[header,["model-B","path-2","C2"],["model-A","path-1","C1"]]
    make(first,a)
    make(second,b,readme_note="changed",illustrative=("beta","alpha"))
    result=review(first,second,required_hashes=False)
    by={s["sheet"]:s for s in result["sheets"]}
    assert by["meta"]["status"]=="IDENTITY_JOIN_IDENTICAL_POSITIONAL_DIFFERENCES_ONLY"
    assert by["meta"]["identity_rows_with_different_row_index"]==2
    assert by["meta"]["row_signatures_only_official"]==0
    assert by["meta_Ch3vetted_withclimate"]["identity_join_identical"] is True
    assert by["Illustrative_Pathways"]["status"]=="SAME_NORMALIZED_ROWS_POSITIONAL_DIFFERENCES_ONLY"
    assert by["README"]["status"]=="NON_IDENTITY_SHEET_CONTENT_DIFFERENCE_UNRESOLVED"
    assert "model-A" not in str(result)
    assert "changed" not in str(result)


def test_hash_gate_blocks_wrong_workbooks(tmp_path):
    one=tmp_path/"a.xlsx";two=tmp_path/"b.xlsx"
    make(one,[["Model","Scenario","Category"],["m","s","C1"]])
    make(two,[["Model","Scenario","Category"],["m","s","C1"]])
    with pytest.raises(SheetReviewError,match="hash"):
        review(one,two,required_hashes=True)


def test_identity_content_change_detected(tmp_path):
    one=tmp_path/"a.xlsx";two=tmp_path/"b.xlsx"
    make(one,[["Model","Scenario","Category"],["m","s","C1"]])
    make(two,[["Model","Scenario","Category"],["m","s","C2"]])
    result=review(one,two,required_hashes=False)
    record=next(x for x in result["sheets"] if x["sheet"]=="meta")
    assert record["status"]=="IDENTITY_JOIN_CONTENT_DIFFERENCE"
