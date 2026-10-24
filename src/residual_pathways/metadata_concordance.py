"""Compare two AR6 metadata workbooks without exporting scenario identities or values.

Outputs normalized cell-level and Model+Scenario-level difference *counts*,
sheet names, source byte hashes, and explicit unresolved source provenance.
This is not a publisher-authentication test.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal
from hashlib import sha256
import json
import math
from pathlib import Path
import unicodedata

from .stratify import sha_file

SCHEMA = "ar6-metadata-concordance-public.v1"


class ConcordanceError(ValueError):
    pass


def normalized(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return ["boolean", value]
    if isinstance(value, (int, float, Decimal)):
        if not math.isfinite(float(value)):
            raise ConcordanceError("Nonfinite metadata cell value")
        return ["number", format(Decimal(str(value)).normalize(), "f")]
    if isinstance(value, (datetime, date, time)):
        return ["date_time", value.isoformat()]
    if isinstance(value, str):
        return ["text", unicodedata.normalize("NFC", value.replace("\r\n", "\n").replace("\r", "\n"))]
    raise ConcordanceError(f"Unsupported metadata cell type {type(value).__name__}")


def normalized_text(value):
    return normalized(value) == ["text", ""] or normalized(value) is None


def _load(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ConcordanceError("Install the optional metadata dependencies: pip install -e '.[metadata]'") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = {}
    for sheet in workbook:
        cells = {}
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
            for col_index, cell in enumerate(row, 1):
                v = normalized(cell)
                if v is not None and v != ["text", ""]:
                    cells[(row_index, col_index)] = v
        result[sheet.title] = cells
    workbook.close()
    return result


def _records(cells):
    row_1 = {c: v for (r, c), v in cells.items() if r == 1}
    header = {tuple(v): c for c, v in row_1.items()}
    model_col = header.get(("text", "Model"))
    scenario_col = header.get(("text", "Scenario"))
    if model_col is None or scenario_col is None:
        return None
    if len(header) != len(row_1):
        raise ConcordanceError("Duplicate metadata headers")
    rows = {}
    row_columns = {}
    for (r, col) in cells:
        if r > 1:
            row_columns.setdefault(r, set()).add(col)
    for r in sorted(row_columns):
        model = cells.get((r, model_col))
        scenario = cells.get((r, scenario_col))
        if not model or not scenario:
            if any(col not in (model_col, scenario_col) for col in row_columns[r]):
                raise ConcordanceError("Metadata data row missing Model or Scenario")
            continue
        key = (json.dumps(model, ensure_ascii=False), json.dumps(scenario, ensure_ascii=False))
        if key in rows:
            raise ConcordanceError("Duplicate Model+Scenario metadata identity")
        rows[key] = {tuple(row_1[c]): cells.get((r, c)) for c in row_1}
    return rows


def compare(official_path, previous_path):
    official_path, previous_path = Path(official_path), Path(previous_path)
    official = _load(official_path)
    previous = _load(previous_path)
    sheets = []
    for name in sorted(set(official) | set(previous)):
        a = official.get(name, {})
        b = previous.get(name, {})
        different_coordinates = sum(a.get(k) != b.get(k) for k in set(a) | set(b))
        item = {
            "sheet": name, "in_official_named": name in official,
            "in_earlier_workbook": name in previous,
            "nonblank_cells_official": len(a),
            "nonblank_cells_earlier": len(b),
            "normalized_cell_differences": different_coordinates,
        }
        rec_a, rec_b = _records(a) if name in official else None, _records(b) if name in previous else None
        if rec_a is not None and rec_b is not None:
            keys_a, keys_b = set(rec_a), set(rec_b)
            shared = keys_a & keys_b
            changed = sum(rec_a[key] != rec_b[key] for key in shared)
            changes = Counter()
            for key in shared:
                for col in set(rec_a[key]) | set(rec_b[key]):
                    if rec_a[key].get(col) != rec_b[key].get(col):
                        changes[str(col[1])] += 1
            item["identity_join"] = {
                "model_scenarios_official": len(keys_a),
                "model_scenarios_earlier": len(keys_b),
                "shared_model_scenarios": len(shared),
                "official_only": len(keys_a - keys_b),
                "earlier_only": len(keys_b - keys_a),
                "shared_records_with_differences": changed,
                "changed_column_counts": dict(sorted(changes.items())),
            }
        sheets.append(item)
    return {
        "schema_version": SCHEMA,
        "provenance_status": "LOCAL_WORKBOOK_COMPARISON_NOT_PUBLISHER_AUTHENTICATION",
        "official_named": {"filename": official_path.name, "sha256": sha_file(official_path)},
        "earlier_redistribution": {"filename": previous_path.name, "sha256": sha_file(previous_path)},
        "status": "EQUIVALENT_NORMALIZED_CELLS" if all(s["normalized_cell_differences"] == 0 and s["in_official_named"] and s["in_earlier_workbook"] for s in sheets) else "DIFFERENCES_RECORDED",
        "sheets": sheets,
        "disclosure": "Only aggregate comparison statistics, sheet/column names and file hashes; no model/scenario names or cell values.",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--official", type=Path, required=True)
    parser.add_argument("--earlier", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path("results/ar6-canonical-gate/metadata_concordance_public.json"))
    args = parser.parse_args()
    result = compare(args.official, args.earlier)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    print(result["status"])
    for sheet in result["sheets"]:
        print(sheet["sheet"], "normalized cell differences:", sheet["normalized_cell_differences"])


if __name__ == "__main__":
    main()
