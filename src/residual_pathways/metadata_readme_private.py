"""Private exact-cell examination of the one unresolved AR6 metadata README difference.

Two workbook hashes are pinned; output is permitted only inside ignored
data/ar6-private. Never commit/display full original README cell contents.
No source provenance authentication or scenario semantics inferred.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import json

from .metadata_concordance import normalized
from .stratify import sha_file

OFFICIAL_SHA = "fd1bc32db8463769e18f0560ced603a49598c12e1abfa06787ff763c55f06716"
EARLIER_SHA = "7d22453bdda52e95cea4df46d362018951dd7114e62f01f5bcea65d6e24a7790"


class ReadmeAuditError(ValueError):
    pass


def _cells(file):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ReadmeAuditError("Install openpyxl using pip install -e '.[metadata]'") from exc
    wb = load_workbook(file, read_only=True, data_only=True)
    try:
        if "README" not in wb:
            raise ReadmeAuditError("README worksheet not present")
        sheet = wb["README"]
        out = {}
        for ri, row in enumerate(sheet.iter_rows(values_only=True), 1):
            for ci, value in enumerate(row, 1):
                nv=normalized(value)
                if nv is not None and nv != ["text",""]:
                    out[ri,ci] = {"type":nv[0], "value":nv[1]}
        return out
    finally:
        wb.close()


def compare_readme(official, earlier, *, pinned=True):
    a,b=Path(official),Path(earlier)
    if pinned and (sha_file(a)!=OFFICIAL_SHA or sha_file(b)!=EARLIER_SHA):
        raise ReadmeAuditError("Workbook bytes differ from registered local sources")
    aa,bb=_cells(a),_cells(b)
    changes=[]
    for key in sorted(set(aa)|set(bb)):
        av,bv=aa.get(key),bb.get(key)
        if av!=bv:
            changes.append({"row":key[0],"column":key[1],
                            "official_value":av,"earlier_value":bv,
                            "reviewer_classification":"UNREVIEWED"})
    return {
        "schema":"ar6-metadata-readme-private-diff.v1",
        "handling":"LOCAL_ONLY_DO_NOT_COMMIT_DO_NOT_UPLOAD",
        "workbook_authentication":"NOT_ESTABLISHED",
        "changes":changes,
        "number_of_changes":len(changes),
        "notice":"Inspect human-readable original text locally; avoid pasting account URLs/tokens. A README difference alone does not alter the fully identity-matched scenario metadata sheets.",
    }


def write(official, earlier, output):
    output=Path(output).resolve()
    if "ar6-private" not in output.parts:
        raise ReadmeAuditError("Output must be private and gitignored under data/ar6-private")
    result=compare_readme(official,earlier)
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(json.dumps(result,indent=2,sort_keys=True,ensure_ascii=False)+"\n",encoding="utf-8")
    return result


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--official",type=Path,required=True)
    p.add_argument("--earlier",type=Path,required=True)
    p.add_argument("--output",type=Path,default=Path("data/ar6-private/metadata_readme_diff_private.json"))
    a=p.parse_args()
    out=write(a.official,a.earlier,a.output)
    print("README changed normalized cells:",out["number_of_changes"])
    for row in out["changes"]:
        print(f"Cell row={row['row']} col={row['column']}: PRIVATE_REVIEW_REQUIRED")
    print("Contents retained ONLY locally:",a.output)


if __name__=="__main__":
    main()
